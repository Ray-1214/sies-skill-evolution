"""
p5_0_probe.py — P5-0 可行性探針
================================
問一個從沒被測過的問題：**本機 Gemma-26B + 這個沙箱，做不做得出簡報 / 遊戲 /
資料管線？**整個 W16 計畫（P3 的邊、P5 的叢集、demo 的開場、3D 圖的兩個區塊）
都押在這個假設上。

9 題 × 3 領域，不走課程系統，直接丟 agent_runner.run()。

⚠️ TD-11：沙箱容器長生命期、任務間無隔離。實測 /tmp 已有 49 項前幾次跑的殘留
   （data.csv / input.csv / sales.csv / output.json …），正是資料管線題會產生的
   檔名。**不清的話 verifier 會因殘留誤判 pass，膨脹通過率，導致 §3.4 做出錯誤
   的領域決策**——那正是這支探針要避免的事。所以每題之間清 /tmp。

輸出 data/p5_0_probe/<ts>.json，供 §3.4.1 的領域決策與 §2.2.1 的去重下限判斷。

用法:
    python scripts/p5_0_probe.py               # 全跑
    python scripts/p5_0_probe.py --only pipeline   # 只跑一個領域
    python scripts/p5_0_probe.py --dry-run     # 只印題目與 verifier，不跑
"""

import argparse
import json
import subprocess
import sys
import time
from datetime import datetime, timezone, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

CONTAINER = "sies-agent-zero"
PY = "/opt/venv/bin/python3"
TZ = timezone(timedelta(hours=8))
OUT_DIR = ROOT / "data" / "p5_0_probe"

# /tmp 底下這些不能清（容器自己的東西）
TMP_KEEP = {"models", "pulse-PKdhtXMmr18n"}


# ---------------------------------------------------------------------------
# 沙箱 helper
# ---------------------------------------------------------------------------

def sh(cmd: str, timeout: int = 120):
    """在容器裡跑一段 bash，回傳 (exit_code, stdout, stderr)。"""
    r = subprocess.run(
        ["docker", "exec", "-e", "SDL_VIDEODRIVER=dummy", CONTAINER, "bash", "-lc", cmd],
        capture_output=True, text=True, timeout=timeout,
    )
    return r.returncode, r.stdout, r.stderr


def clean_tmp():
    """TD-11：把 /tmp 歸零（保留容器自己的目錄）。回傳清掉幾項。"""
    keep = " ".join(f"! -name '{k}'" for k in TMP_KEEP)
    code, out, _ = sh(f"ls -A /tmp | wc -l")
    before = int(out.strip() or 0)
    sh(f"find /tmp -mindepth 1 -maxdepth 1 {keep} -exec rm -rf {{}} + 2>/dev/null; true")
    code, out, _ = sh("ls -A /tmp | wc -l")
    after = int(out.strip() or 0)
    return before, after


def put_file(path: str, content: str):
    """把內容寫進容器裡的檔案（base64 避開 escaping）。"""
    import base64
    b = base64.b64encode(content.encode()).decode()
    sh(f"mkdir -p $(dirname {path}) && echo '{b}' | base64 -d > {path}")


# ---------------------------------------------------------------------------
# 9 題 + verifier
# ---------------------------------------------------------------------------
# oracle 獨立性（§7.1）：
#   pipeline → verified_independent（我先放已知輸入，verifier 比對已知正解）
#   presentation / game → verified_liveness（只驗開得起來 / 跑得動 / 結構合法）

SALES_CSV = """region,product,units,price
North,widget,10,2.5
North,gadget,,4.0
South,widget,20,2.5
South,gadget,5,4.0
East,widget,,2.5
East,gadget,8,4.0
"""

TASKS = [
    # ---------------- pipeline（獨立 oracle）----------------
    dict(
        tid="D1", domain="pipeline", cluster="C2", level=1,
        setup=lambda w: put_file(f"{w}/sales.csv", SALES_CSV),
        task=(
            f"Read the CSV file at WORKDIR/sales.csv. The 'units' column has missing values. "
            f"Fill each missing 'units' value with the mean of the non-missing 'units' values, "
            f"rounded to 2 decimal places. Write the cleaned data to WORKDIR/clean.csv with the "
            f"same columns and the same row order. Then print exactly one line: "
            f"ROWS=<number of data rows> FILLED=<number of values you filled>."
        ),
        # 已知正解：units = 10,?,20,5,?,8 → 非缺值平均 (10+20+5+8)/4 = 10.75，填 2 格，6 列
        verify=r"""
import csv,sys
p="WORKDIR/clean.csv"
rows=list(csv.DictReader(open(p)))
assert len(rows)==6, f"row count {len(rows)} != 6"
u=[float(r["units"]) for r in rows]
exp=[10.0,10.75,20.0,5.0,10.75,8.0]
for i,(a,b) in enumerate(zip(u,exp)):
    assert abs(a-b)<0.011, f"row {i}: {a} != {b}"
print("VERIFY_PASS")
""",
        oracle="verified_independent",
    ),
    dict(
        tid="D2", domain="pipeline", cluster="C1", level=2,
        setup=lambda w: put_file(f"{w}/records.json", json.dumps([
            {"id": 1, "name": "alice", "age": 30},
            {"id": 2, "name": "bob"},
            {"id": "x", "name": "carol", "age": 25},
            {"id": 4, "name": "dave", "age": -1},
        ])),
        task=(
            "Read WORKDIR/records.json (a JSON list of objects). Write a schema validator that "
            "checks each record satisfies: 'id' is an int, 'name' is a non-empty string, 'age' is "
            "an int >= 0 and is present. Write the validation result to WORKDIR/result.json as a "
            "JSON list of booleans, one per record, in the same order. Then print exactly one "
            "line: VALID=<number of valid records>."
        ),
        # 已知正解：[True, False(缺 age), False(id 非 int), False(age<0)] → VALID=1
        verify=r"""
import json
r=json.load(open("WORKDIR/result.json"))
assert r==[True,False,False,False], f"got {r}, expected [True,False,False,False]"
print("VERIFY_PASS")
""",
        oracle="verified_independent",
    ),
    dict(
        tid="D3", domain="pipeline", cluster="C3", level=2,
        setup=lambda w: put_file(f"{w}/sales.csv", SALES_CSV),
        task=(
            "Read WORKDIR/sales.csv. Ignore rows where 'units' is empty. Compute total revenue "
            "(units * price) grouped by 'region'. Write WORKDIR/report.json as a JSON object "
            "mapping region name to total revenue (a number). Then print exactly one line: "
            "REGIONS=<number of regions in your output>."
        ),
        # 已知正解：North widget 10*2.5=25（gadget 缺值忽略）；South 20*2.5+5*4=70；East 8*4=32
        verify=r"""
import json
d=json.load(open("WORKDIR/report.json"))
exp={"North":25.0,"South":70.0,"East":32.0}
assert set(d)==set(exp), f"regions {sorted(d)} != {sorted(exp)}"
for k,v in exp.items():
    assert abs(float(d[k])-v)<0.01, f"{k}: {d[k]} != {v}"
print("VERIFY_PASS")
""",
        oracle="verified_independent",
    ),

    # ---------------- presentation（liveness）----------------
    dict(
        tid="P1", domain="presentation", cluster="C2", level=1,
        setup=None,
        task=(
            "Using the python-pptx library, create a PowerPoint file at WORKDIR/deck.pptx with "
            "exactly 3 slides: slide 1 is a title slide with the title 'Q3 Sales Review'; slides 2 "
            "and 3 each have a title and a bulleted list of at least 3 bullet points about sales "
            "performance. Then print exactly one line: SLIDES=<number of slides>."
        ),
        verify=r"""
from pptx import Presentation
p=Presentation("WORKDIR/deck.pptx")
assert len(p.slides)==3, f"{len(p.slides)} slides != 3"
texts=[[sh.text_frame.text for sh in s.shapes if sh.has_text_frame] for s in p.slides]
assert any("Q3 Sales Review" in t for t in texts[0]), "slide 1 missing title"
for i in (1,2):
    body=max(texts[i], key=len) if texts[i] else ""
    n=len([l for l in body.split("\n") if l.strip()])
    assert n>=3, f"slide {i+1} has {n} bullet lines, need >=3"
print("VERIFY_PASS")
""",
        oracle="verified_liveness",
    ),
    dict(
        tid="P2", domain="presentation", cluster="C2", level=2,
        setup=lambda w: put_file(f"{w}/sales.csv", SALES_CSV),
        task=(
            "Read WORKDIR/sales.csv. For each distinct 'region', create one slide in a PowerPoint "
            "file at WORKDIR/regions.pptx. Each slide's title must be exactly the region name, and "
            "its body must list the products sold in that region. Use python-pptx. Then print "
            "exactly one line: SLIDES=<number of slides>."
        ),
        verify=r"""
from pptx import Presentation
p=Presentation("WORKDIR/regions.pptx")
titles={s.shapes.title.text.strip() for s in p.slides if s.shapes.title}
assert titles=={"North","South","East"}, f"titles {sorted(titles)}"
print("VERIFY_PASS")
""",
        oracle="verified_liveness",
    ),
    dict(
        tid="P3", domain="presentation", cluster="C5", level=2,
        setup=lambda w: put_file(f"{w}/para.txt",
            "Our third quarter results exceeded expectations across every region. "
            "Widget sales grew steadily in the north while gadget adoption accelerated "
            "in the south. Supply chain delays that hampered the second quarter were "
            "largely resolved by August. Customer retention improved by four points. "
            "We expect momentum to continue into the fourth quarter."),
        task=(
            "Read the paragraph in WORKDIR/para.txt. Compress it into at most 5 bullet points, "
            "each at most 12 words. Put them on a single slide in a PowerPoint file at "
            "WORKDIR/summary.pptx, with the slide title 'Q3 Summary'. Use python-pptx. Then print "
            "exactly one line: BULLETS=<number of bullets>."
        ),
        verify=r"""
from pptx import Presentation
p=Presentation("WORKDIR/summary.pptx")
assert len(p.slides)==1, f"{len(p.slides)} slides != 1"
s=p.slides[0]
body=""
for sh in s.shapes:
    if sh.has_text_frame and (not s.shapes.title or sh!=s.shapes.title):
        if len(sh.text_frame.text)>len(body): body=sh.text_frame.text
lines=[l.strip() for l in body.split("\n") if l.strip()]
assert 1<=len(lines)<=5, f"{len(lines)} bullets, need 1..5"
bad=[l for l in lines if len(l.split())>12]
assert not bad, f"bullets over 12 words: {bad}"
print("VERIFY_PASS")
""",
        oracle="verified_liveness",
    ),

    # ---------------- game（liveness）----------------
    dict(
        tid="G1", domain="game", cluster="C1", level=1,
        setup=None,
        task=(
            "Write a headless pygame program at WORKDIR/bounce.py that creates a 320x240 display, "
            "simulates one ball bouncing off all four walls for exactly 120 frames (no real-time "
            "delay, no event loop blocking), and on the last frame prints exactly one line: "
            "POS=<x>,<y> with the integer ball position. The ball must always stay within the "
            "window bounds. Run it and show the output. Set SDL_VIDEODRIVER=dummy before importing "
            "pygame so it runs without a display."
        ),
        verify=r"""
import subprocess,re,os
env=dict(os.environ); env["SDL_VIDEODRIVER"]="dummy"
r=subprocess.run(["/opt/venv/bin/python3","WORKDIR/bounce.py"],capture_output=True,text=True,timeout=60,env=env)
assert r.returncode==0, f"exit {r.returncode}: {r.stderr[-300:]}"
m=re.search(r"POS=(-?\d+),(-?\d+)", r.stdout)
assert m, f"no POS= line in output: {r.stdout[-300:]}"
x,y=int(m.group(1)),int(m.group(2))
assert 0<=x<=320 and 0<=y<=240, f"ball out of bounds: {x},{y}"
print("VERIFY_PASS")
""",
        oracle="verified_liveness",
    ),
    dict(
        tid="G2", domain="game", cluster="C1", level=2,
        setup=None,
        task=(
            "Write a headless pygame program at WORKDIR/collide.py. Create two pygame.Rect objects: "
            "A starts at (0,100) size 20x20 moving right 5 px/frame; B starts at (200,100) size "
            "20x20 moving left 5 px/frame. Advance frame by frame. On the first frame where they "
            "overlap, print exactly one line: COLLISION=<frame number> and stop. Set "
            "SDL_VIDEODRIVER=dummy. Run it and show the output."
        ),
        verify=r"""
import subprocess,re,os
env=dict(os.environ); env["SDL_VIDEODRIVER"]="dummy"
r=subprocess.run(["/opt/venv/bin/python3","WORKDIR/collide.py"],capture_output=True,text=True,timeout=60,env=env)
assert r.returncode==0, f"exit {r.returncode}: {r.stderr[-300:]}"
m=re.search(r"COLLISION=(\d+)", r.stdout)
assert m, f"no COLLISION= line: {r.stdout[-300:]}"
f=int(m.group(1))
assert 15<=f<=20, f"collision at frame {f}, expected ~18"
print("VERIFY_PASS")
""",
        oracle="verified_liveness",
    ),
    dict(
        tid="G3", domain="game", cluster="C3", level=2,
        setup=None,
        task=(
            "Write a headless pygame program at WORKDIR/bullets.py simulating a simple bullet-hell "
            "tick. Spawn exactly 5 bullets at the centre of a 320x240 screen, each with a different "
            "constant velocity. Advance exactly 60 frames with no real-time delay. Then print "
            "exactly one line: OFFSCREEN=<how many of the 5 bullets are now outside the 320x240 "
            "screen>. Set SDL_VIDEODRIVER=dummy. Run it and show the output."
        ),
        verify=r"""
import subprocess,re,os
env=dict(os.environ); env["SDL_VIDEODRIVER"]="dummy"
r=subprocess.run(["/opt/venv/bin/python3","WORKDIR/bullets.py"],capture_output=True,text=True,timeout=60,env=env)
assert r.returncode==0, f"exit {r.returncode}: {r.stderr[-300:]}"
m=re.search(r"OFFSCREEN=(\d+)", r.stdout)
assert m, f"no OFFSCREEN= line: {r.stdout[-300:]}"
n=int(m.group(1))
assert 0<=n<=5, f"OFFSCREEN={n} not in 0..5"
print("VERIFY_PASS")
""",
        oracle="verified_liveness",
    ),
]


# ---------------------------------------------------------------------------
# 遊戲題 v2：句式對齊簡報題
# ---------------------------------------------------------------------------
# v1 寫成「Write a program AT <path> ... Run it and show the output」，把 agent
# 推向 bash + 寫檔路徑，它做完 mkdir 就捏造了剩下的（3/3 全部 fabricate）。
# v2 改成「用 Python 寫出檔案 → 執行它 → print 一行」，跟 P1-P3 同一個句式。
# 檔案仍然要存在（verifier 不變），差別只在怎麼把 agent 導向 code_execution。

GAME_V2 = [
    dict(
        tid="G1b", domain="game", cluster="C1", level=1, setup=None,
        oracle="verified_liveness",
        task=(
            "Write a Python program to the file WORKDIR/bounce.py using Python's open(). "
            "The program must set os.environ['SDL_VIDEODRIVER']='dummy' before importing pygame, "
            "create a 320x240 display, simulate one ball bouncing off all four walls for exactly "
            "120 frames with no real-time delay, and print exactly one line POS=<x>,<y> with the "
            "integer ball position, always within the window bounds. After writing the file, "
            "execute it with subprocess and print exactly one line: RESULT=<the POS line it printed>."
        ),
        verify=[t for t in TASKS if t["tid"] == "G1"][0]["verify"].replace("bounce.py", "bounce.py"),
    ),
    dict(
        tid="G2b", domain="game", cluster="C1", level=2, setup=None,
        oracle="verified_liveness",
        task=(
            "Write a Python program to the file WORKDIR/collide.py using Python's open(). "
            "The program must set os.environ['SDL_VIDEODRIVER']='dummy' before importing pygame, "
            "then create two pygame.Rect objects: A at (0,100) size 20x20 moving right 5 px/frame, "
            "B at (200,100) size 20x20 moving left 5 px/frame. Advance frame by frame and on the "
            "first frame where they overlap print exactly one line COLLISION=<frame number> and "
            "stop. After writing the file, execute it with subprocess and print exactly one line: "
            "RESULT=<the COLLISION line it printed>."
        ),
        verify=[t for t in TASKS if t["tid"] == "G2"][0]["verify"],
    ),
    dict(
        tid="G3b", domain="game", cluster="C3", level=2, setup=None,
        oracle="verified_liveness",
        task=(
            "Write a Python program to the file WORKDIR/bullets.py using Python's open(). "
            "The program must set os.environ['SDL_VIDEODRIVER']='dummy' before importing pygame, "
            "spawn exactly 5 bullets at the centre of a 320x240 screen each with a different "
            "constant velocity, advance exactly 60 frames with no real-time delay, and print "
            "exactly one line OFFSCREEN=<how many of the 5 are outside the screen>. After writing "
            "the file, execute it with subprocess and print exactly one line: RESULT=<that line>."
        ),
        verify=[t for t in TASKS if t["tid"] == "G3"][0]["verify"],
    ),
]


def run_verifier(t, workdir):
    """在沙箱裡跑 verifier。回傳 (passed, detail)。"""
    code = t["verify"].replace("WORKDIR", workdir)
    vpath = f"{workdir}/_verify.py"
    put_file(vpath, code)
    rc, out, err = sh(f"cd {workdir} && {PY} {vpath}", timeout=120)
    if rc == 0 and "VERIFY_PASS" in out:
        return True, "ok"
    tail = (err or out).strip().splitlines()
    return False, (tail[-1][:220] if tail else f"exit {rc}")


def main():
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument("--only", choices=["pipeline", "presentation", "game"])
    ap.add_argument("--game-v2", action="store_true",
                    help="只跑重寫過句式的遊戲題（分辨『不會 pygame』vs『題目誘發捏造』）")
    ap.add_argument("--dry-run", action="store_true")
    args = ap.parse_args()

    tasks = GAME_V2 if args.game_v2 else [
        t for t in TASKS if not args.only or t["domain"] == args.only]

    if args.dry_run:
        for t in tasks:
            print(f"\n=== {t['tid']} [{t['domain']}/{t['cluster']}/L{t['level']}] "
                  f"oracle={t['oracle']} ===")
            print(t["task"].replace("WORKDIR", f"/tmp/p50_{t['tid']}"))
        return 0

    from agents.agent_runner import AgentRunner
    runner = AgentRunner("config.yaml")

    OUT_DIR.mkdir(parents=True, exist_ok=True)
    ts = datetime.now(TZ).strftime("%Y%m%dT%H%M%S")
    results = []

    print(f"[P5-0] {len(tasks)} 題，每題之間清 /tmp（TD-11）\n")

    for i, t in enumerate(tasks, 1):
        wd = f"/tmp/p50_{t['tid']}"

        # ── TD-11：跑之前把 /tmp 歸零 ──
        before, after = clean_tmp()
        sh(f"mkdir -p {wd}")
        if t["setup"]:
            t["setup"](wd)

        prompt = t["task"].replace("WORKDIR", wd)
        print(f"[{i}/{len(tasks)}] {t['tid']} [{t['domain']}/{t['cluster']}] "
              f"(/tmp 清掉 {before}→{after} 項)")
        print(f"        {prompt[:100]}...")

        t0 = time.time()
        rec = dict(tid=t["tid"], domain=t["domain"], cluster=t["cluster"],
                   level=t["level"], oracle=t["oracle"], workdir=wd)
        try:
            r = runner.run(prompt)
            rec.update(
                execution_completed=bool(r.execution_success),
                steps=r.execution_steps,
                error_stage=r.error_stage,
                error=(str(r.error)[:200] if r.error else None),
                candidates=[c.name for c in r.evaluation.validated_candidates]
                           if r.evaluation else [],
            )
        except Exception as e:  # noqa: BLE001
            rec.update(execution_completed=False, steps=0,
                       error_stage="runner_exception", error=f"{type(e).__name__}: {e}",
                       candidates=[])

        passed, detail = run_verifier(t, wd)
        rec.update(verified_success=passed, verify_detail=detail,
                   elapsed_seconds=round(time.time() - t0, 1))
        # 留下產物清單當證據
        _, ls, _ = sh(f"ls -A {wd} 2>/dev/null")
        rec["artifacts"] = [x for x in ls.split() if not x.startswith("_verify")]

        mark = "✓ PASS" if passed else "✗ FAIL"
        print(f"        exec={rec['execution_completed']} steps={rec['steps']} "
              f"{rec['elapsed_seconds']}s  verify: {mark} ({detail})")
        print(f"        產物: {rec['artifacts']}\n")
        results.append(rec)

    out = OUT_DIR / f"{ts}.json"
    out.write_text(json.dumps({"generated": ts, "results": results},
                              ensure_ascii=False, indent=2), encoding="utf-8")

    # ── 摘要 ──
    print("=" * 68)
    print("P5-0 結果")
    print("=" * 68)
    by = {}
    for r in results:
        d = by.setdefault(r["domain"], {"n": 0, "exec": 0, "ver": 0, "steps": []})
        d["n"] += 1
        d["exec"] += bool(r["execution_completed"])
        d["ver"] += bool(r["verified_success"])
        d["steps"].append(r["steps"])
    print(f"{'領域':14s} {'verified':>10s} {'exec_done':>10s} {'avg steps':>10s}")
    for dom, d in by.items():
        avg = sum(d['steps']) / d['n'] if d['n'] else 0
        print(f"{dom:14s} {d['ver']}/{d['n']:>8} {d['exec']}/{d['n']:>8} {avg:>10.1f}")
    print(f"\n→ {out}")
    print("\n§3.4.1 決策規則：簡報 ≥2/3 且 遊戲 ≥2/3 → 40/30/30；"
          "簡報 ≥2/3 遊戲 ≤1/3 → 簡報 60/pipeline 20/general 20；"
          "簡報 ≤1/3 → pipeline 60/general 40")
    return 0


if __name__ == "__main__":
    sys.exit(main())
